from datetime import datetime
import os
from openpyxl import load_workbook, Workbook
import streamlit as st

# Cấu hình trang web cho di động
st.set_page_config(
    page_title="TRADING BOT - NHẬP LỆNH", page_icon="📈", layout="centered"
)


# --- LOGIC GHI EXCEL NÂNG CẤP (HỖ TRỢ FILE EXE) ---
def ghi_du_lieu_trading(loai, entry, sl, tp, ma, label_ket_qua):
    # KIỂM TRA: Nếu đang chạy dạng file .exe, lấy đường dẫn của file .exe đó
    if getattr(sys, "frozen", False):
        thu_muc_hien_tai = os.path.dirname(sys.executable)
    # Nếu chạy dạng file .py thông thường
    else:
        thu_muc_hien_tai = os.path.dirname(os.path.abspath(__file__))

    # Khóa chặt đường dẫn file Excel tại thư mục chứa app/code
    file_name = os.path.join(thu_muc_hien_tai, "trading_TC.xlsx")

    # Xử lý mặc định mã là XAUUSD nếu để trống
    ma = ma.strip() if ma else "XAUUSD"
    if not ma:
        ma = "XAUUSD"

    if not loai or not entry or not sl or not tp:
        label_ket_qua.configure(
            text="Lỗi: Vui lòng nhập đầy đủ thông tin!", fg_color="#FF3333"
        )
        return

    try:
        if os.path.exists(file_name):
            wb = load_workbook(file_name)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws["B1"], ws["D1"], ws["E1"], ws["F1"], ws["J1"], ws["K1"] = (
                "Loại",
                "Entry",
                "SL",
                "TP",
                "Mã",
                "Thời gian tín hiệu",
            )

        # Tìm dòng trống thực tế
        next_row = 2
        while True:
            if (
                ws[f"B{next_row}"].value is None
                and ws[f"D{next_row}"].value is None
                and ws[f"E{next_row}"].value is None
                and ws[f"F{next_row}"].value is None
                and ws[f"J{next_row}"].value is None
            ):
                break
            next_row += 1

        thoi_gian_hien_tai = datetime.now().strftime("%d/%m/%Y %H:%M")

        # Ghi dữ liệu
        ws[f"B{next_row}"] = loai
        ws[f"D{next_row}"] = float(entry) if entry.replace(".", "", 1).isdigit() else entry
        ws[f"E{next_row}"] = float(sl) if sl.replace(".", "", 1).isdigit() else sl
        ws[f"F{next_row}"] = float(tp) if tp.replace(".", "", 1).isdigit() else tp
        ws[f"J{next_row}"] = ma.upper()
        ws[f"K{next_row}"] = thoi_gian_hien_tai

        wb.save(file_name)

        label_ket_qua.configure(
            text=f"Thành công: Đã ghi lệnh {loai.upper()} {ma.upper()} vào dòng {next_row}!",
            fg_color="#2EB73A",
        )

    except Exception as e:
        label_ket_qua.configure(
            text=f"Lỗi khi lưu file: {str(e)}", fg_color="#FF3333"
        )


# --- GIAO DIỆN WEB UI (TƯƠNG THÍCH ĐIỆN THOẠI) ---
st.markdown(
    "<h2 style='text-align: center;'>TRADING TC ENTRY</h2>",
    unsafe_allow_html=True,
)
st.markdown(
    "<p style='text-align: center; color: gray;'>Quản lý lệnh từ xa trên điện thoại</p>",
    unsafe_allow_html=True,
)

with st.form("trading_form"):
    # 1. Mã Giao Dịch[cite: 2]
    ma = st.text_input(
        "Mã Sản Phẩm (BTCUSDT, XAUUSD...):",
        placeholder="Nhập mã...(Mặc định:XAUUSD)",
    )

    # 2. Loại lệnh (Buy / Sell)[cite: 2]
    loai = st.radio("Loại Lệnh:", options=["buy", "sell"], horizontal=True)

    # 3, 4, 5. Entry, SL, TP[cite: 2]
    entry = st.text_input("Giá Vào Lệnh (Entry):", placeholder="Ví dụ: 2350")
    sl = st.text_input("Cắt Lỗ (SL):", placeholder="Ví dụ: 2340")
    tp = st.text_input("Chốt Lời (TP):", placeholder="Ví dụ: 2370")

    # Nút bấm gửi lệnh
    submit_button = st.form_submit_button(
        label="GỬI LỆNH VÀO EXCEL", use_container_width=True
    )

    if submit_button:
        success, message = ghi_du_lieu_trading(loai, entry, sl, tp, ma)
        if success:
            st.success(message)
        else:
            st.error(message)